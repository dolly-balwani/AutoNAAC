"""
Excel Parser for NAAC Criterion Data
Extracts events with their Google Docs proof links
"""
import pandas as pd
from openpyxl import load_workbook
from typing import List, Dict
import re

def parse_criterion_sheet(excel_path: str, sheet_name: str) -> List[Dict]:
    """
    Parse a criterion sheet from the Excel file.
    Returns a list of events with their details including hyperlinks.
    """
    try:
        # Load workbook to extract hyperlinks
        wb = load_workbook(excel_path)
        ws = wb[sheet_name]
        
        # Build hyperlink map (cell -> URL)
        hyperlinks = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.hyperlink:
                    hyperlinks[(cell.row, cell.column)] = cell.hyperlink.target
        
        # Read with pandas for easier data handling
        df = pd.read_excel(excel_path, sheet_name=sheet_name, header=1)
        df.columns = [str(col).strip().replace('\n', ' ')[:50] for col in df.columns]
        
        events = []
        current_category = None
        
        for idx, row in df.iterrows():
            excel_row = idx + 3  # Account for header row offset
            
            first_col = str(row.iloc[0]) if pd.notna(row.iloc[0]) else ""
            
            # Check if this is a category header
            if re.match(r'^\d+\.\s+\w+', first_col) and pd.isna(row.iloc[1]):
                current_category = first_col
                continue
            
            if first_col == "" or first_col == "nan":
                continue
            
            # Get hyperlink from the Document Proof column (usually column 5 = E)
            doc_link = hyperlinks.get((excel_row, 5), "")
            
            event = {
                "category": current_category,
                "name": first_col,
                "date": str(row.iloc[1]) if len(row) > 1 and pd.notna(row.iloc[1]) else "",
                "students": str(row.iloc[2]) if len(row) > 2 and pd.notna(row.iloc[2]) else "",
                "agencies": str(row.iloc[3]) if len(row) > 3 and pd.notna(row.iloc[3]) else "",
                "doc_link": doc_link
            }
            
            if event["name"] and (event["date"] or event["students"] or event["doc_link"]):
                events.append(event)
        
        wb.close()
        return events
    
    except Exception as e:
        print(f"Error parsing sheet {sheet_name}: {e}")
        return []

def get_available_criteria(excel_path: str) -> List[str]:
    """Get list of criterion sheets available in the Excel file."""
    try:
        xl = pd.ExcelFile(excel_path)
        # Filter to only criterion-named sheets (e.g., 5.1.3, 5.2.2)
        criteria = [s for s in xl.sheet_names if re.match(r'^\d+\.\d+\.\d+', s)]
        return criteria
    except Exception as e:
        print(f"Error reading Excel: {e}")
        return []

if __name__ == "__main__":
    excel_path = "data/Criteria 5.1.3 CMPN Data 2024-25.xlsx"
    
    print("Available criteria sheets:")
    criteria = get_available_criteria(excel_path)
    print(criteria)
    
    print("\nParsing 5.1.3...")
    events = parse_criterion_sheet(excel_path, "5.1.3")
    print(f"Found {len(events)} events")
    
    for i, event in enumerate(events[:5]):
        print(f"\n{i+1}. {event['name']}")
        print(f"   Date: {event['date']}")
        print(f"   Students: {event['students']}")
        print(f"   Doc Link: {event['doc_link'][:50]}..." if event['doc_link'] else "   Doc Link: None")
